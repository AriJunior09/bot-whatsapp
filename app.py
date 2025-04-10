"""
AutoBot - Automação de Mensagens
"""
import openpyxl
from urllib.parse import quote
import webbrowser
from time import sleep
import pyautogui
import os
from datetime import datetime
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import threading

# Variáveis globais
caminho_planilha = None
mensagem_padrao = "Olá {nome}, seu boleto vence no dia {vencimento}. Favor pagar no link https://www.link_do_pagamento.com"
pausado = False

def carregar_planilha():
    global caminho_planilha
    caminho_planilha = filedialog.askopenfilename(filetypes=[("Arquivos Excel", "*.xlsx")])
    if caminho_planilha:
        label_status.config(text=f"Planilha carregada: {caminho_planilha}")
        messagebox.showinfo("Sucesso", "Planilha carregada com sucesso!")
    else:
        label_status.config(text="Nenhuma planilha carregada.")

def alternar_pausa():
    """Alterna o estado de pausa."""
    global pausado
    pausado = not pausado
    if pausado:
        btn_pausar.config(text="Retomar")
        label_status.config(text="Envio pausado.")
    else:
        btn_pausar.config(text="Pausar")
        label_status.config(text="Envio retomado.")

def iniciar_envio():
    global caminho_planilha, mensagem_padrao, pausado
    if not caminho_planilha:
        messagebox.showerror("Erro", "Nenhuma planilha carregada. Por favor, carregue uma planilha antes de iniciar o envio.")
        return

    # Obter a mensagem personalizada do campo de texto
    mensagem_personalizada = campo_mensagem.get("1.0", tk.END).strip()
    if not mensagem_personalizada:
        messagebox.showerror("Erro", "A mensagem personalizada não pode estar vazia.")
        return

    try:
        # Abrir a planilha
        workbook = openpyxl.load_workbook(caminho_planilha)
        pagina_clientes = workbook.active

        # Variável para rastrear se houve erros
        erros = False

        # Loop pelos dados da planilha
        for linha in pagina_clientes.iter_rows(min_row=2):
            while pausado:
                label_status.config(text="Envio pausado. Aguarde...")
                sleep(1)  # Aguarda enquanto o envio está pausado

            nome = linha[0].value
            telefone = linha[1].value
            vencimento = linha[2].value

            # Validar os dados
            if not nome or not telefone or not vencimento:
                print(f"⚠️ Dados incompletos para {nome}. Pulando...")
                erros = True
                continue

            # Formatar a data de vencimento
            if isinstance(vencimento, str):
                try:
                    vencimento = datetime.strptime(vencimento, "%d/%m/%Y")
                except ValueError:
                    print(f"⚠️ Data inválida para {nome}. Pulando...")
                    erros = True
                    continue

            # Substituir os placeholders na mensagem personalizada
            mensagem = mensagem_personalizada.format(
                nome=nome,
                vencimento=vencimento.strftime('%d/%m/%Y')
            )

            # Abrir o WhatsApp Web e enviar a mensagem
            try:
                link_mensagem_whatsapp = f"https://web.whatsapp.com/send?phone={telefone}&text={quote(mensagem)}"
                webbrowser.open(link_mensagem_whatsapp)
                sleep(15)
                pyautogui.hotkey('enter')  # Pressionar Enter para enviar
                sleep(5)
                pyautogui.hotkey('ctrl', 'w')  # Fechar a aba
                sleep(5)
            except Exception as e:
                print(f"❌ Erro ao enviar mensagem para {nome}: {e}")
                erros = True
                with open('erros.csv', 'a', newline='', encoding='utf-8') as arquivo:
                    arquivo.write(f"{nome},{telefone}\n")

        # Exibir mensagem de conclusão
        if erros:
            messagebox.showwarning("Concluído com Erros", "Envio concluído, mas alguns contatos apresentaram erros. Verifique o relatório de erros.")
        else:
            messagebox.showinfo("Concluído", "Todas as mensagens foram enviadas com sucesso!")

    except Exception as e:
        messagebox.showerror("Erro", f"Ocorreu um erro ao processar a planilha: {e}")

def iniciar_envio_thread():
    """Inicia o envio em uma thread separada."""
    thread_envio = threading.Thread(target=iniciar_envio)
    thread_envio.daemon = True  # Permite que o programa feche mesmo com a thread em execução
    thread_envio.start()

def exibir_relatorio():
    if os.path.exists('erros.csv'):
        os.startfile('erros.csv')
    else:
        messagebox.showinfo("Relatório", "Nenhum erro registrado até o momento.")

# Criar janela principal
janela = tk.Tk()
janela.title("AutoBot - Automação de Mensagens")
janela.geometry("700x600")
janela.resizable(False, False)

# Estilo
style = ttk.Style()
style.theme_use("clam")  # Tema moderno
style.configure("TButton", font=("Arial", 12), padding=6)
style.configure("TLabel", font=("Arial", 12))

# Cabeçalho
frame_cabecalho = ttk.Frame(janela, padding=10)
frame_cabecalho.pack(fill="x")

label_titulo = ttk.Label(frame_cabecalho, text="AutoBot - Automação de Mensagens", font=("Arial", 16, "bold"))
label_titulo.pack()

# Área de configuração
frame_config = ttk.LabelFrame(janela, text="Configuração", padding=10)
frame_config.pack(fill="x", padx=10, pady=10)

btn_carregar = ttk.Button(frame_config, text="Carregar Planilha", command=carregar_planilha)
btn_carregar.pack(side="left", padx=5)

label_status = ttk.Label(frame_config, text="Nenhuma planilha carregada.")
label_status.pack(side="left", padx=10)

# Área de mensagem personalizada
frame_mensagem = ttk.LabelFrame(janela, text="Mensagem Personalizada", padding=10)
frame_mensagem.pack(fill="both", expand=True, padx=10, pady=10)

campo_mensagem = tk.Text(frame_mensagem, height=5, font=("Arial", 12))
campo_mensagem.pack(fill="both", expand=True)
campo_mensagem.insert("1.0", mensagem_padrao)  # Inserir mensagem padrão

# Área de execução
frame_execucao = ttk.LabelFrame(janela, text="Execução", padding=10)
frame_execucao.pack(fill="x", padx=10, pady=10)

btn_iniciar = ttk.Button(frame_execucao, text="Iniciar Envio", command=iniciar_envio_thread)
btn_iniciar.pack(side="left", padx=5)

btn_pausar = ttk.Button(frame_execucao, text="Pausar", command=alternar_pausa)
btn_pausar.pack(side="left", padx=5)

btn_relatorio = ttk.Button(frame_execucao, text="Exibir Relatório de Erros", command=exibir_relatorio)
btn_relatorio.pack(side="left", padx=5)

btn_sair = ttk.Button(frame_execucao, text="Sair", command=janela.quit)
btn_sair.pack(side="left", padx=5)

# Área de status
frame_status = ttk.LabelFrame(janela, text="Status", padding=10)
frame_status.pack(fill="x", padx=10, pady=10)

label_status = ttk.Label(frame_status, text="Aguardando ação do usuário.")
label_status.pack()

# Rodapé
frame_rodape = ttk.Frame(janela, padding=10)
frame_rodape.pack(fill="x", side="bottom")

label_rodape = ttk.Label(frame_rodape, text="Desenvolvido por Ari Júnior", font=("Arial", 10, "italic"))
label_rodape.pack(side="right", padx=10)

# Iniciar o loop da interface
janela.mainloop()
